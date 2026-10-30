"""Preserve Korean document text, filenames and existing Word font metadata."""
from io import BytesIO
import json
import zipfile

from docx import Document
from docx.oxml.ns import qn

from backend.documents import remote_image_docx


def test_docx_preserves_korean_text_and_existing_font():
    original = Document()
    run = original.add_paragraph().add_run('기밀 테스트 문서 — 홍길동')
    run.font.name = '맑은 고딕'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    buffer = BytesIO()
    original.save(buffer)
    result = remote_image_docx(buffer.getvalue(), 'https://test.example/b/test-token', 'test-file')
    document = Document(BytesIO(result))
    preserved = document.paragraphs[0].runs[0]
    assert preserved.text == '기밀 테스트 문서 — 홍길동'
    assert preserved.font.name == '맑은 고딕'
    assert preserved._element.rPr.rFonts.get(qn('w:eastAsia')) == '맑은 고딕'


def test_zip_preserves_korean_filename_and_utf8_content(system):
    client, headers = system['client'], system['headers']
    text = '승인된 가상 테스트 자료입니다. 원본 보존을 확인합니다.'
    registered = client.post('/api/files', headers=headers,
        files={'file': ('기밀 문서.txt', text.encode('utf-8'))}, data={'classification': 'CONFIDENTIAL'})
    registered.raise_for_status()
    created = client.post('/api/distributions', headers=headers,
        json={'file_ids': [registered.json()['id']], 'approved_test_documents': True})
    created.raise_for_status()
    download = client.get('/api/distributions/' + created.json()['id'] + '/download', headers=headers)
    download.raise_for_status()
    with zipfile.ZipFile(BytesIO(download.content)) as archive:
        item = json.loads(archive.read('manifest.json'))['files'][0]
        assert item['file_name'] == '기밀 문서.txt'
        assert '기밀 문서' in item['archive_path']
        assert archive.read(item['archive_path']).decode('utf-8').startswith(text)
