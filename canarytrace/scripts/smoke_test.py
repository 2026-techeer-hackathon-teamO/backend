"""Isolated loopback HTTPS integration; fixture device reports are NOT physical PC evidence.

Creates a separate runtime under ignored artifacts and never connects to the live server.
No Slack, email or external AI calls. Does not change Windows audit/network settings.
"""
import hashlib
import io
import json
import os
from pathlib import Path
import socket
import ssl
import subprocess
import sys
import time
import uuid
import zipfile
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import httpx
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from dotenv import dotenv_values
from agent.windows_agent import Agent
from scripts import setup


def main():
    work = ROOT / 'artifacts' / ('smoke-' + str(uuid.uuid4()))
    work.mkdir(parents=True)
    with socket.socket() as port_probe:
        port_probe.bind(('127.0.0.1', 0))
        port = port_probe.getsockname()[1]
    old_root, old_argv = setup.ROOT, sys.argv
    try:
        setup.ROOT, sys.argv = work, ['setup', '--port', str(port)]
        setup.main()
    finally:
        setup.ROOT, sys.argv = old_root, old_argv
    runtime = work / '.runtime'
    credentials = dotenv_values(work / '.env')
    env = {**os.environ, **credentials, 'DATAGUARD_RUNTIME': str(runtime),
           'SLACK_WEBHOOK_URL': '', 'PYTHONUTF8': '1', 'CORPORATE_NETWORKS': '127.0.0.0/8,::1/128'}
    base_url = f'https://localhost:{port}'
    ca = runtime / 'tls/cert.pem'
    checks = {}
    process = None

    def launch(log):
        child = subprocess.Popen([sys.executable, '-X', 'utf8', '-m', 'backend.run'], cwd=ROOT,
            env=env, stdout=log, stderr=log, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        return child

    def ready(client):
        for _ in range(100):
            if process.poll() is not None:
                raise RuntimeError('Isolated server exited; inspect ' + str(work / 'server.log'))
            try:
                if client.get('/api/health').status_code == 200:
                    return
            except httpx.TransportError:
                pass
            time.sleep(.2)
        raise RuntimeError('Isolated HTTPS server did not become ready')

    def login(client):
        r = client.post('/api/login', json={'username': 'admin', 'password': credentials['DATAGUARD_ADMIN_PASSWORD']})
        r.raise_for_status()
        return {'Authorization': 'Bearer ' + r.json()['access_token']}

    try:
        with (work / 'server.log').open('w', encoding='utf-8') as log, httpx.Client(
                base_url=base_url, verify=ssl.create_default_context(cafile=str(ca)), trust_env=False, timeout=10) as client:
            process = launch(log)
            ready(client)
            checks['real_loopback_https'] = True
            headers = login(client)
            original = Document()
            original.add_paragraph('한글 가상 테스트 문서입니다. 실제 기밀자료가 아닙니다.')
            buffer = io.BytesIO()
            original.save(buffer)
            r = client.post('/api/files', headers=headers, files={'file': ('한글 테스트.docx', buffer.getvalue())}, data={'classification': 'CONFIDENTIAL'})
            r.raise_for_status()
            r = client.post('/api/distributions', headers=headers, json={'file_ids': [r.json()['id']], 'approved_test_documents': True, 'docx_remote_image': True})
            r.raise_for_status()
            distribution = r.json()
            path = '/api/distributions/' + distribution['id']
            assert client.get(path + '/download').status_code == 401
            checks['unauthenticated_download_rejected'] = True
            downloaded = client.get(path + '/download', headers=headers)
            downloaded.raise_for_status()
            assert hashlib.sha256(downloaded.content).hexdigest() == distribution['sha256']
            with zipfile.ZipFile(io.BytesIO(downloaded.content)) as archive:
                assert archive.testzip() is None
                item = json.loads(archive.read('manifest.json'))['files'][0]
                document_bytes = archive.read(item['archive_path'])
                assert hashlib.sha256(document_bytes).hexdigest() == item['file_hash']
            checks['zip_download_and_member_hashes'] = True
            doc = Document(io.BytesIO(document_bytes))
            url = next(rel.target_ref for rel in doc.part.rels.values() if rel.is_external and rel.reltype == RT.IMAGE)
            with client.stream('GET', '/api/stream', headers=headers) as stream:
                stream.raise_for_status()
                lines = stream.iter_lines()
                before = next(line for line in lines if line.startswith('data:'))
                beacon = client.get(url)
                assert beacon.status_code == 200 and beacon.content.startswith(b'\x89PNG')
                after = next(line for line in lines if line.startswith('data:'))
                assert before != after
            checks['real_beacon_http_and_sse_refresh'] = True
            r = client.post('/api/devices', headers=headers, json={'name': 'SYNTHETIC-TRANSPORT-TEST'})
            r.raise_for_status()
            device = r.json()
            approved = work / 'approved'
            approved.mkdir()
            config = work / 'device.json'
            config.write_text(json.dumps({'server': base_url, 'device_id': device['id'], 'device_key': device['device_key'], 'ca_cert': str(ca), 'roots': [str(approved)]}), 'utf-8')
            agent = Agent(config)
            try:
                agent.refresh()
                body = {'file_id': item['file_id'], 'sha256': item['file_hash'], 'record_id': 1,
                    'occurred_at': datetime.now(timezone.utc).isoformat(), 'pc_name': 'SYNTHETIC-TRANSPORT-TEST',
                    'windows_user': '가상테스트\\홍길동', 'windows_sid': 'SYNTHETIC-SID',
                    'process': 'C:\\Synthetic\\notepad.exe', 'process_id': '0x1234',
                    'object_path': str(approved / '가상문서.docx'), 'access_mask': 1,
                    'source_type': 'WINDOWS_SECURITY_4663', 'storage_type': 'LOCAL'}
                agent.enqueue('/api/agent/events', body, 'synthetic-transport-test')
                agent.db.commit()
                agent.flush()
                assert agent.db.execute('SELECT COUNT(*) FROM pending').fetchone()[0] == 0
            finally:
                agent.close()
            events = client.get('/api/events', headers=headers).json()
            reported = next(e for e in events if e.get('device_id') == device['id'])
            assert reported['issuance_id'] == item['issuance_id']
            assert reported['windows_user'] == '가상테스트\\홍길동'
            checks['signed_agent_transport_with_synthetic_payload'] = True
            event_ids = {e['event_id'] for e in events}
            process.terminate()
            process.wait(timeout=15)
            process = launch(log)
            ready(client)
            headers = login(client)
            assert event_ids <= {e['event_id'] for e in client.get('/api/events', headers=headers).json()}
            assert client.get('/api/integrity', headers=headers).json()['ok']
            checks['restart_persistence_and_integrity'] = True
    finally:
        if process is not None and process.poll() is None:
            process.terminate()
            process.wait(timeout=15)
        report = {'checks': checks, 'status': 'PASS' if len(checks) == 6 and all(checks.values()) else 'FAIL',
                  'physical_pc_usb_email': 'NOT_TESTED', 'device_payload': 'SYNTHETIC_ISOLATED_FIXTURE'}
        (work / 'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), 'utf-8')
        print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report['status'] == 'PASS' else 1


if __name__ == '__main__':
    raise SystemExit(main())
