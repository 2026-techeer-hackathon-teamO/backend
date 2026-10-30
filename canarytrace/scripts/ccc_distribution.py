"""Inspect the actual Windows Desktop/ccc and create a local approved ZIP via HTTPS."""
import argparse
import hashlib
import io
import json
import os
import re
import ssl
import subprocess
import sys
import uuid
import zipfile
from contextlib import contextmanager
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import httpx
from docx import Document
from dotenv import dotenv_values
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from backend.archives import source_files, zip_entries
from backend.documents import validate_document


def desktop_context():
    if os.name != 'nt':
        raise RuntimeError('Windows Desktop discovery required')
    result = subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-Command',
        "[Console]::OutputEncoding=[Text.UTF8Encoding]::new(); [PSCustomObject]@{user=[Environment]::UserName;profile=[Environment]::GetFolderPath('UserProfile');desktop=[Environment]::GetFolderPath('DesktopDirectory');registry_desktop=(Get-ItemProperty -LiteralPath 'HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\Explorer\\User Shell Folders').Desktop;onedrive=$env:OneDrive}|ConvertTo-Json -Compress"],
        check=True, capture_output=True, encoding='utf-8', creationflags=subprocess.CREATE_NO_WINDOW)
    context = json.loads(result.stdout.lstrip('\ufeff'))
    desktop = Path(context['desktop'])
    if not desktop.is_dir():
        raise RuntimeError('Actual Desktop is unavailable')
    context['candidates'] = [str(p) for p in desktop.iterdir() if p.name.casefold() == 'ccc' or p.name.casefold().startswith('ccc.')]
    return context


def file_info(name, raw):
    info = {'name': name, 'extension': Path(name).suffix.lower(), 'size': len(raw), 'sha256': hashlib.sha256(raw).hexdigest(), 'readable': True, 'is_zip': zipfile.is_zipfile(io.BytesIO(raw)), 'encrypted': False, 'actual_format': 'UNKNOWN', 'existing_canary': 'NOT_DETECTED_HEURISTIC', 'supported': False}
    scan = [raw]
    try:
        if info['is_zip']:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                info['encrypted'] = any(i.flag_bits & 1 for i in archive.infolist())
                info['members'] = [{'name': i.filename, 'size': i.file_size, 'compressed_size': i.compress_size, 'encrypted': bool(i.flag_bits & 1)} for i in archive.infolist()]
                info['actual_format'] = 'DOCX' if 'word/document.xml' in archive.namelist() else 'ZIP'
            scan = list(zip_entries(raw).values())
        elif raw.startswith(b'%PDF-'):
            info['actual_format'] = 'PDF'
            pdf = PdfReader(io.BytesIO(raw))
            info['encrypted'] = pdf.is_encrypted
            if not pdf.is_encrypted:
                scan.append(str(pdf.trailer).encode())
                for page in pdf.pages:
                    scan.append((page.extract_text() or '').encode())
                    scan.append(str([a.get_object() for a in page.get('/Annots', [])]).encode())
        else:
            raw.decode('utf-8-sig')
            if b'\x00' not in raw:
                info['actual_format'] = 'TXT_UTF8'
        pattern = rb'(?i)(canarytokens?\.|canarytrace|/api/canary/|https?://[^\s<>"\x27]+/(?:t|b)/[A-Za-z0-9_-]{20,})'
        if any(re.search(pattern, part) for part in scan):
            info['existing_canary'] = 'POSSIBLE_EXISTING_TOKEN'
        validate_document(name, raw)
        info['supported'] = True
    except Exception as exc:
        info['validation_error'] = type(exc).__name__ + ': ' + str(exc)[:200]
    return info


def inspect_source(source):
    source = Path(source)
    content = source_files(source)
    files = []
    for name, raw in content.items():
        info = file_info(name, raw)
        path = source / name if source.is_dir() else source
        info['filesystem_encrypted'] = bool(getattr(path.stat(), 'st_file_attributes', 0) & 0x4000)
        files.append(info)
    return {'source_path': str(source), 'kind': 'DIRECTORY' if source.is_dir() else 'FILE', 'extension': source.suffix if source.is_file() else None, 'filesystem_encrypted': bool(getattr(source.stat(), 'st_file_attributes', 0) & 0x4000), 'readable': True, 'file_count': len(content), 'files': files, 'canary_scan_limitation': 'Heuristic only; absence of a known marker does not prove absence of all tracking.', 'fingerprint': {n: hashlib.sha256(b).hexdigest() for n, b in content.items()}}, content


def synthetic_documents():
    docs = {'ccc-test.txt': b'APPROVED SYNTHETIC TEST ONLY. No business data.\nOpening TXT alone sends no HTTP request.\n'}
    buffer = io.BytesIO()
    doc = Document(); doc.add_heading('CCC approved synthetic security test', 0)
    doc.add_paragraph('Fictional demonstration data only. No confidential source content.'); doc.save(buffer)
    docs['ccc-test.docx'] = buffer.getvalue()
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer); pdf.drawString(40, 760, 'CCC approved SYNTHETIC security test - no business data'); pdf.save()
    docs['ccc-test.pdf'] = buffer.getvalue()
    return docs


@contextmanager
def api_client():
    runtime = ROOT / '.runtime'
    config = json.loads((runtime / 'settings.json').read_text('utf-8'))
    c = httpx.Client(base_url=config['public_base_url'], verify=ssl.create_default_context(cafile=str(runtime / 'tls/cert.pem')), trust_env=False, timeout=60)
    try:
        response = c.post('/api/login', json={'username': 'admin', 'password': dotenv_values(ROOT / '.env')['DATAGUARD_ADMIN_PASSWORD']})
        response.raise_for_status()
        c.headers['Authorization'] = 'Bearer ' + response.json()['access_token']
        yield c
    finally:
        if 'Authorization' in c.headers:
            c.post('/api/logout')
        c.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inspect-only', action='store_true')
    parser.add_argument('--synthetic', action='store_true', help='Use fictional fixtures outside the original ccc')
    parser.add_argument('--approved-test-files', action='store_true', help='Explicit operator approval that all selected source files are test data')
    parser.add_argument('--source', type=Path, help='Resolve ambiguity between discovered ccc candidates')
    parser.add_argument('--beacon-html', action='store_true', help='Add declared HTML remote-image test copies')
    parser.add_argument('--docx-remote-image', action='store_true', help='Explicitly approve a visible external image reference in DOCX copies')
    args = parser.parse_args()
    context = desktop_context()
    if args.source:
        source = args.source.absolute()
        if str(source) not in context['candidates']:
            raise SystemExit('--source must be one of the actual Desktop ccc candidates')
    elif len(context['candidates']) == 1:
        source = Path(context['candidates'][0])
    else:
        raise SystemExit('Expected one ccc candidate; found: ' + json.dumps(context['candidates']))
    report, content = inspect_source(source)
    report['desktop_context'] = context
    out = ROOT / 'artifacts' / 'distributions' / str(uuid.uuid4())
    out.mkdir(parents=True, exist_ok=False)
    (out / 'source-inspection.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), 'utf-8')
    print(json.dumps({'source': str(source), 'kind': report['kind'], 'files': report['file_count'], 'inspection': str(out / 'source-inspection.json')}, ensure_ascii=False), flush=True)
    if args.inspect_only:
        return
    if args.synthetic:
        docs = synthetic_documents()
        fixture = out / 'synthetic-source'; fixture.mkdir()
        for name, raw in docs.items():
            (fixture / name).write_bytes(raw)
    else:
        if not args.approved_test_files:
            raise SystemExit('Explicit --approved-test-files required; use --synthetic for confidential or unknown material')
        docs = {}
        for name, raw in content.items():
            info = file_info(name, raw)
            if info['actual_format'] == 'ZIP':
                for child, data in zip_entries(raw).items():
                    member = name + '/' + child
                    docs[member] = data
            else:
                docs[name] = raw
        if not docs:
            raise SystemExit('ccc is empty. Use --synthetic; original stays unchanged.')
        for name, raw in docs.items():
            info = file_info(name, raw)
            if not info['supported'] or info['existing_canary'] != 'NOT_DETECTED_HEURISTIC':
                raise SystemExit('Unsupported, encrypted, or previously tracked source; inspect report and use --synthetic: ' + name)
    with api_client() as c:
        registered = c.get('/api/files'); registered.raise_for_status()
        known = registered.json(); file_ids = []
        for name, raw in docs.items():
            sha = hashlib.sha256(raw).hexdigest()
            match = next((f for f in known if f['sha256'] == sha), None)
            if match is None:
                r = c.post('/api/files', files={'file': (Path(name).name, raw)}, data={'classification': 'CONFIDENTIAL'})
                r.raise_for_status(); match = r.json(); known.append(match)
            file_ids.append(match['id'])
        r = c.post('/api/distributions', json={'name': 'ccc', 'source_label': 'ccc (synthetic replacement)' if args.synthetic else 'ccc', 'source_kind': 'SYNTHETIC_TEST_DATA' if args.synthetic else 'APPROVED_SOURCE', 'approved_test_documents': True, 'file_ids': list(dict.fromkeys(file_ids)), 'include_beacon_html': args.beacon_html, 'docx_remote_image': args.docx_remote_image})
        r.raise_for_status(); distribution = r.json()
        downloaded = c.get('/api/distributions/' + distribution['id'] + '/download'); downloaded.raise_for_status()
        if hashlib.sha256(downloaded.content).hexdigest() != distribution['sha256']:
            raise RuntimeError('Download hash mismatch')
        entries = zip_entries(downloaded.content)
        manifest = json.loads(entries['manifest.json'])
        for item in manifest['files']:
            if hashlib.sha256(entries[item['archive_path']]).hexdigest() != item['file_hash']:
                raise RuntimeError('ZIP member hash mismatch')
        (out / distribution['filename']).write_bytes(downloaded.content)
        (out / 'manifest.json').write_bytes(entries['manifest.json'])
        (out / 'distribution-id.txt').write_text(distribution['id'] + '\n', 'utf-8')
        (out / 'token-metadata.json').write_text(json.dumps([{key: item.get(key) for key in ['distribution_id', 'issuance_id', 'file_id', 'decoy_id', 'source_decoy_id', 'token_id', 'file_hash', 'mode', 'created_at']} for item in manifest['files']], ensure_ascii=False, indent=2), 'utf-8')
        (out / 'distribution.json').write_text(json.dumps(distribution, ensure_ascii=False, indent=2), 'utf-8')
        after, _ = inspect_source(source)
        if after['fingerprint'] != report['fingerprint']:
            raise RuntimeError('Original source changed during processing')
        result = {'distribution_id': distribution['id'], 'zip_path': str(out / distribution['filename']), 'source_preserved': True, 'synthetic': args.synthetic, 'zip_hash_verified': True, 'member_hashes_verified': True, 'source_path': str(source), 'physical_remote_pc': 'NOT_TESTED', 'access_status': 'UNKNOWN'}
        (out / 'creation-results.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), 'utf-8')
        (ROOT / 'artifacts' / 'distributions' / 'latest.json').write_text(json.dumps({'output': str(out), **result}, ensure_ascii=False, indent=2), 'utf-8')
        print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
