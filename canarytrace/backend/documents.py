"""Declared links/optional DOCX image: no macros, execution or security bypasses."""
from io import BytesIO
import zipfile
import html
import struct
import zlib
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

LIMIT = 10 * 1024 * 1024


def tracking_png():
    """Visible blue indicator, not a hidden tracking pixel."""
    def chunk(kind, data):
        return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!2I5B', 48, 24, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress((b'\x00' + b'\x23\x73\xb5' * 48) * 24)) + chunk(b'IEND', b'')


def remote_image_docx(content, url, file_id):
    from docx.shared import Inches
    from urllib.parse import urlparse
    if urlparse(url).scheme != 'https':
        raise ValueError('DOCX image requires HTTPS')
    doc = Document(BytesIO(traced_copy(content, 'docx', url, file_id)))
    doc.add_heading('APPROVED CANARYTRACE REMOTE IMAGE', level=2)
    doc.add_paragraph('This visible test image requests the security server only if Word permits external images. '
                      'The server records time, IP and User-Agent. A request does not prove human reading. '
                      'If blocked, access remains unobserved. Do not change Protected View or security policy for this test.')
    shape = doc.add_picture(BytesIO(tracking_png()), width=Inches(1), height=Inches(.5))
    blip = shape._inline.xpath('.//a:blip')[0]
    blip.attrib.pop(qn('r:embed'), None)
    blip.set(qn('r:link'), doc.part.relate_to(url, RT.IMAGE, is_external=True))
    shape._inline.docPr.set('descr', 'Approved CanaryTrace remote image; network request only when allowed by Word')
    output = BytesIO(); doc.save(output)
    return output.getvalue()


def beacon_document(name, url, file_id):
    return ('<!doctype html><html lang="ko"><meta charset="utf-8"><title>Approved Canary test</title>'
            '<h1>Approved Canary test document</h1><p>' + html.escape(name) + '</p><p>File ID: ' + html.escape(file_id) + '</p>'
            '<p>This approved test requests a remote image when online and permitted by your viewer. '
            'The server records request time and IP. It cannot identify a PC or prove human reading from IP.</p>'
            '<img alt="Canary test beacon" width="1" height="1" src="' + html.escape(url, quote=True) + '"></html>').encode('utf-8')


def validate_document(name: str, content: bytes):
    ext = name.rsplit('.', 1)[-1].lower()
    if ext not in {'txt', 'docx', 'pdf'}:
        raise ValueError('Only TXT, DOCX, PDF are supported')
    if not content or len(content) > LIMIT:
        raise ValueError('File must be nonempty and at most 10 MiB')
    if ext == 'txt':
        content.decode('utf-8-sig')
    if ext == 'docx':
        with zipfile.ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024 or len(z.infolist()) > 2000:
                raise ValueError('Document expansion limit exceeded')
            if any('vbaproject' in n.lower() for n in z.namelist()):
                raise ValueError('Macro-enabled documents are not accepted')
        Document(BytesIO(content))
    if ext == 'pdf':
        pdf = PdfReader(BytesIO(content))
        if pdf.is_encrypted:
            raise ValueError('Encrypted PDF is unsupported')
        if len(pdf.pages) > 500:
            raise ValueError('PDF page limit exceeded')
    return ext


def traced_copy(content: bytes, ext: str, url: str, file_id: str) -> bytes:
    label = 'Enterprise security tracking - approved CanaryTrace verification link (URL access only)'
    if ext == 'txt':
        return content + f'\n\n[{label}]\nFile ID: {file_id}\n{url}\nOpening this TXT alone sends no network request.\n'.encode()
    if ext == 'docx':
        doc = Document(BytesIO(content))
        doc.add_paragraph(f'{label}\nFile ID: {file_id}')
        p = doc.add_paragraph()
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = url
        r.append(t)
        link.append(r)
        p._p.append(link)
        out = BytesIO()
        doc.save(out)
        return out.getvalue()
    source = PdfReader(BytesIO(content))
    extra = BytesIO()
    c = canvas.Canvas(extra)
    c.drawString(40, 790, label)
    c.drawString(40, 765, f'File ID: {file_id}')
    c.setFont('Helvetica', 8)
    c.drawString(40, 740, url)
    c.linkURL(url, (40, 730, 550, 755), relative=0)
    c.drawString(40, 710, 'Opening this PDF alone is not tracked. No scripts or embedded code added.')
    c.save()
    writer = PdfWriter()
    writer.append(source)
    writer.append(PdfReader(extra))
    out = BytesIO()
    writer.write(out)
    return out.getvalue()
